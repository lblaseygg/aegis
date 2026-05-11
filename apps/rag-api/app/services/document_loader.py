from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

from app.utils.file_types import file_type, is_supported


@dataclass(slots=True)
class LoadedDocument:
    path: Path
    text: str
    file_type: str

    @property
    def file_name(self) -> str:
        return self.path.name


class DocumentLoader:
    def iter_paths(self, root: Path, recursive: bool = True) -> list[Path]:
        if root.is_file():
            return [root] if is_supported(root) else []

        iterator = root.rglob("*") if recursive else root.glob("*")
        return [path for path in iterator if path.is_file() and is_supported(path)]

    def load(self, path: Path) -> LoadedDocument:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".csv"}:
            text = path.read_text(encoding="utf-8", errors="ignore")
        elif suffix in {".html", ".htm"}:
            text = self._load_html(path)
        elif suffix == ".json":
            payload = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
            text = json.dumps(payload, indent=2, sort_keys=True)
        elif suffix == ".docx":
            text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
        elif suffix == ".pdf":
            reader = PdfReader(str(path))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            raise ValueError(f"Unsupported file type for {path}")

        return LoadedDocument(path=path, text=text.strip(), file_type=file_type(path))

    def _load_html(self, path: Path) -> str:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
        return soup.get_text(separator="\n")
